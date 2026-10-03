import openpyxl
import json
import os
import datetime

excel_dir = r"C:\Users\shena\OneDrive\Desktop\NATILLERA 2026"
out_dir = r"C:\Users\shena\natillera-2026"

def serialize_date(d):
    if isinstance(d, (datetime.datetime, datetime.date)):
        return d.strftime('%Y-%m-%d')
    return str(d) if d is not None else ''

# 1. PARSE AHORROS
wb_ahorro = openpyxl.load_workbook(os.path.join(excel_dir, "Control_Natillera_Ahorro.xlsx"), data_only=True)
ws_res = wb_ahorro['Resumen General']

ahorradores = []
for r in range(9, 22):
    num = ws_res.cell(r, 1).value
    nombre = ws_res.cell(r, 2).value
    if nombre:
        modalidad = ws_res.cell(r, 3).value
        cuota = ws_res.cell(r, 4).value
        meta = ws_res.cell(r, 5).value
        ahorrado = ws_res.cell(r, 6).value
        ahorradores.append({
            "id": int(num) if num else len(ahorradores)+1,
            "nombre": str(nombre).strip(),
            "modalidad": str(modalidad).strip() if modalidad else "Quincenal",
            "cuotaBase": float(cuota) if cuota else 40000,
            "metaAnual": float(meta) if meta else 960000,
            "telefono": ""
        })

ws_mat = wb_ahorro['Matriz Quincenal']
member_cols = {}
for c in range(4, ws_mat.max_column + 1, 2):
    m_name = ws_mat.cell(6, c).value
    if m_name:
        member_cols[str(m_name).strip()] = (c, c + 1)

quincenas = []
for r in range(8, 32):
    fecha_raw = ws_mat.cell(r, 1).value
    q_code = ws_mat.cell(r, 2).value
    if not q_code:
        continue
    pagos = {}
    for m_name, (c_val, c_forma) in member_cols.items():
        v = ws_mat.cell(r, c_val).value
        f = ws_mat.cell(r, c_forma).value
        pagos[m_name] = {
            "valor": float(v) if v is not None and v != '' else 0,
            "formaPago": str(f).strip() if f is not None else ""
        }
    quincenas.append({
        "quincena": str(q_code).strip(),
        "fecha": str(fecha_raw).strip() if fecha_raw else '',
        "pagos": pagos
    })

# 2. PARSE CREDITOS
wb_cred = openpyxl.load_workbook(os.path.join(excel_dir, "Plantilla_Control_Creditos_Natillera_Dashboard.xlsx"), data_only=True)
socios_creditos = ['Jorge Gaviria', 'Caro Rios', 'Andres Castro', 'Wilson Alvarez', 'Ana Lopez']
creditos = []
abonos = []

cred_id_counter = 1
abono_id_counter = 1

for s in socios_creditos:
    ws = wb_cred[s]
    # Credits Table
    for r in range(9, ws.max_row + 1):
        num = ws.cell(r, 1).value
        val_cred = ws.cell(r, 6).value
        if val_cred is not None:
            try:
                val_cred_flt = float(val_cred)
                if val_cred_flt > 0:
                    fecha_c = ws.cell(r, 3).value
                    fecha_p = ws.cell(r, 4).value
                    dias = ws.cell(r, 5).value
                    tasa = ws.cell(r, 7).value
                    int_hoy = ws.cell(r, 8).value
                    int_pago = ws.cell(r, 9).value
                    ab_int = ws.cell(r, 10).value
                    sal_int = ws.cell(r, 11).value
                    ab_cap = ws.cell(r, 12).value
                    sal_cap = ws.cell(r, 13).value
                    tot_pagar = ws.cell(r, 14).value
                    estado = ws.cell(r, 15).value

                    c_code = f"CR-{cred_id_counter:03d}"
                    cred_id_counter += 1

                    creditos.append({
                        "id": c_code,
                        "numInterno": int(num) if isinstance(num, (int, float)) else None,
                        "socio": s,
                        "fechaCredito": serialize_date(fecha_c)[:10],
                        "fechaPagoLiquidacion": serialize_date(fecha_p)[:10] if fecha_p else None,
                        "valorCredito": val_cred_flt,
                        "tasaMes": float(tasa) if tasa is not None else 0.005,
                        "interesesCongelados": float(int_pago) if isinstance(int_pago, (int, float)) else None,
                        "estado": str(estado).strip() if estado else ("PAGADO" if (sal_cap == 0 and sal_int == 0) else "VIGENTE")
                    })
            except Exception:
                pass

        # Abonos Table
        val_abono = ws.cell(r, 22).value  # Total abonado (Col V)
        if val_abono is not None:
            try:
                val_abono_flt = float(val_abono)
                if val_abono_flt > 0:
                    f_abono = ws.cell(r, 17).value
                    num_cred = ws.cell(r, 18).value
                    ab_int_val = ws.cell(r, 20).value
                    ab_cap_val = ws.cell(r, 21).value
                    medio = ws.cell(r, 23).value

                    a_code = f"AB-{abono_id_counter:04d}"
                    abono_id_counter += 1

                    abonos.append({
                        "id": a_code,
                        "socio": s,
                        "fechaPago": serialize_date(f_abono)[:10],
                        "numCredito": int(num_cred) if isinstance(num_cred, (int, float)) else 1,
                        "abonoInteres": float(ab_int_val) if isinstance(ab_int_val, (int, float)) else 0,
                        "abonoCapital": float(ab_cap_val) if isinstance(ab_cap_val, (int, float)) else 0,
                        "totalAbonado": val_abono_flt,
                        "medio": str(medio).strip() if medio else "Transferencia",
                        "comprobante": ""
                    })
            except Exception:
                pass

full_data = {
    "config": {
        "nombre": "NATILLERA 2026",
        "administrador": "Santiago Henao",
        "telefonoAdmin": "",
        "ano": 2026,
        "moneda": "COP",
        "tasaMesDefault": 0.005
    },
    "ahorradores": ahorradores,
    "quincenas": quincenas,
    "creditos": creditos,
    "abonos": abonos
}

out_file = os.path.join(out_dir, "initial_data.json")
with open(out_file, "w", encoding="utf-8") as f:
    json.dump(full_data, f, ensure_ascii=False, indent=2)

print(f"OK! Exported {len(ahorradores)} ahorradores, {len(quincenas)} quincenas, {len(creditos)} creditos, {len(abonos)} abonos.")
